"""
Project Screenshots & Working Demonstration Generator (DOCX & PDF)
Module: src/generate_screenshots_document.py
Compiles deployed dashboard screenshots and econometric diagnostic figures with detailed academic captions.
Produces:
- submission_files/03_Petrochemical_Trade_Analytics_Project_Screenshots.docx
- submission_files/03_Petrochemical_Trade_Analytics_Project_Screenshots.pdf
"""

import os
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

SCREENSHOT_ITEMS = [
    {
        "title": "Figure 1: Executive Command Center & Dynamic KPI Metrics (Dark Mode)",
        "path": "outputs/screenshots/01_executive_dashboard_dark.png",
        "description": "Enterprise command center displaying the core 6-card KPI grid: $302.18B Gross Trade Value, $148.44B Outbound Exports, $153.74B Inbound Imports, -$5.29B Net Balance, 334.13 Mt Physical Volume, and 100% usable data audit. Includes annual trade velocity (2015–2024) and product portfolio share."
    },
    {
        "title": "Figure 2: Interactive Global Choropleth Trade Map & Scorecard",
        "path": "outputs/screenshots/02_global_trade_map_dark.png",
        "description": "Geospatial choropleth world map visualizing trade surplus and deficit intensities across the 8 primary reporter economies (China, Germany, India, Japan, Netherlands, South Korea, Saudi Arabia, USA) and bilateral trading partners, supplemented by the geographical balance scorecard."
    },
    {
        "title": "Figure 3A: Olefins vs. Aromatics Cracking Matrix & Price Distribution (Dark Mode)",
        "path": "outputs/screenshots/03_olefins_aromatics_matrix_dark.png",
        "description": "Chemical family comparative analytics contrasting Light Olefins (Ethylene, Propylene, Butadiene; 51.2% market share) against Aromatic Hydrocarbons (Benzene, Toluene; 48.8% market share), paired with benchmark Unit Price (USD/MT) box plot distributions."
    },
    {
        "title": "Figure 3B: Olefins vs. Aromatics Cracking Matrix & Price Distribution (Normal / Light Mode)",
        "path": "outputs/screenshots/03_olefins_aromatics_matrix_light.png",
        "description": "Light mode interface demonstration verifying full dynamic color contrast, slate typography, razor-sharp Plotly axes, and vibrant category markers without washed-out elements."
    },
    {
        "title": "Figure 4: Bilateral Supply Chain Corridors & Cargo Flow Sankey",
        "path": "outputs/screenshots/04_bilateral_sankey_corridors_dark.png",
        "description": "Interactive multi-node Sankey diagram tracking physical bilateral petrochemical cargo flows from major net-exporting hubs (Saudi Arabia, USA, South Korea) into high-demand importing destinations (China, India, Germany)."
    },
    {
        "title": "Figure 5: Econometric Hypothesis Testing & Statistical Inference Lab",
        "path": "outputs/screenshots/05_econometric_hypothesis_lab_dark.png",
        "description": "Real-time econometric inference portal validating structural market shifts: Welch's t-test for COVID-19 structural break (t = -4.858, p = 1.42e-6, Reject H0), One-way ANOVA for inter-commodity price divergence (F = 11.23, p = 2.70e-9), and Pearson price elasticity."
    },
    {
        "title": "Figure 6: Predictive AI & Market Stress-Testing Simulator",
        "path": "outputs/screenshots/06_predictive_ai_simulator.png",
        "description": "Interactive What-If forecasting engine powered by the trained Gradient Boosting model (R² = 0.9543). Evaluates simulated Brent Crude Oil price shocks ($40–$140/bbl) and cargo volume variations to forecast clearing value and implied Trade Unit Value ($/MT)."
    },
    {
        "title": "Figure 7: Enterprise SQL Analytics Terminal (Star Schema)",
        "path": "outputs/screenshots/07_sql_analytics_terminal.png",
        "description": "Embedded SQL terminal executing analytical queries against the local SQLite Star Schema data warehouse (database/petrochemical_trade.db). Features queries with Common Table Expressions (CTEs), ROW_NUMBER, LAG, and DENSE_RANK window functions."
    },
    {
        "title": "Figure 8: College LMS Submission Deliverables Hub",
        "path": "outputs/screenshots/08_lms_deliverables_hub.png",
        "description": "Direct 1-click download center for all required academic submission artifacts: Project Report (PDF/DOCX), Project Synopsis (PDF/DOCX), Complete Code File (PDF/DOCX), Executive Financial Model (XLSX), Cleaned UN Comtrade CSVs, and Power BI DAX Guide."
    },
    {
        "title": "Figure 9: Machine Learning Model Performance — Actual vs. Predicted (R² = 0.9543)",
        "path": "outputs/figures/actual_vs_predicted.png",
        "description": "Regression diagnostics scatter plot showing predicted versus actual petrochemical trade value across test partitions, demonstrating high explanatory power and low residual dispersion."
    },
    {
        "title": "Figure 10: Predictive Feature Importance Ranking",
        "path": "outputs/figures/feature_importance.png",
        "description": "Feature attribution analysis demonstrating that cargo volume (MT), crude oil benchmark proxies, and origin/partner country encodings dominate trade value determination."
    },
    {
        "title": "Figure 11: Econometric Correlation Heatmap",
        "path": "outputs/figures/correlation_heatmap.png",
        "description": "Correlation matrix evaluating dependencies between transaction value, physical volume, unit clearing prices, crude oil proxies, and chronological time indices."
    },
    {
        "title": "Figure 12: Unsupervised K-Means Bilateral Corridor Clustering (PCA Biplot)",
        "path": "outputs/figures/cluster_biplot_pca.png",
        "description": "Principal Component Analysis (PCA) biplot displaying K-Means clustering segmentation of international trade routes into high-volume bulk corridors, niche specialty flows, and re-export centers."
    }
]

def generate_screenshots_docx(output_path):
    print(f"Generating Screenshots DOCX at: {output_path}...")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    t_p = doc.add_paragraph()
    t_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = t_p.add_run("PETROCHEMICAL TRADE ANALYTICS\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(29, 53, 87)

    r2 = t_p.add_run("PROJECT SCREENSHOTS & DEPLOYED SYSTEM DEMONSTRATION\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 173, 181)

    r3 = t_p.add_run("Functional Visual Verification across Streamlit Web UI, SQL Engine, and ML Analytics Modules\n")
    r3.font.name = "Arial"
    r3.font.size = Pt(10)
    r3.font.italic = True
    r3.font.color.rgb = RGBColor(100, 116, 139)

    # Info Table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Curriculum / Program:", "HCL Career Shaper COE & NASSCOM FutureSkills Prime"),
        ("Project Title:", "Petrochemical Trade Analytics & Predictive Intelligence"),
        ("Candidate Name:", "Anshu Kumar"),
        ("Scope:", "Visual Evidence of Working Application & Analytics")
    ]
    for row_idx, (k, v) in enumerate(meta):
        row = info_table.rows[row_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.text = k
        cell_v.text = v
        cell_k.paragraphs[0].runs[0].font.bold = True
        cell_k.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell_v.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    for idx, item in enumerate(SCREENSHOT_ITEMS, start=1):
        path = item["path"]
        if not os.path.exists(path):
            print(f"Skipping missing image: {path}")
            continue

        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r_h = h.add_run(item["title"])
        r_h.font.name = "Arial"
        r_h.font.size = Pt(12)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(29, 53, 87)

        # Image Paragraph
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_after = Pt(4)
        
        # Check aspect ratio
        with PILImage.open(path) as img:
            w, h_px = img.size
            aspect = h_px / w
        
        # Target width 6.2 inches
        target_w = Inches(6.2)
        target_h = Inches(min(3.8, 6.2 * aspect))
        doc.add_picture(path, width=target_w, height=target_h)

        # Caption
        cap_p = doc.add_paragraph()
        cap_p.paragraph_format.space_after = Pt(14)
        cap_p.paragraph_format.line_spacing = 1.15
        r_cap = cap_p.add_run(f"Description & Analytical Function: {item['description']}")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

        if idx % 2 == 0 and idx < len(SCREENSHOT_ITEMS):
            doc.add_page_break()

    doc.save(output_path)
    print(f"Successfully generated Screenshots DOCX: {output_path} ({os.path.getsize(output_path)} bytes)")

def generate_screenshots_pdf(output_path):
    print(f"Generating Screenshots PDF at: {output_path}...")
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    primary_color = colors.HexColor("#1D3557")
    accent_color = colors.HexColor("#00ADB5")
    text_color = colors.HexColor("#0F172A")

    title_style = ParagraphStyle(
        'ScreenTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=primary_color,
        alignment=1,
        spaceAfter=4
    )
    sub_style = ParagraphStyle(
        'ScreenSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=accent_color,
        alignment=1,
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        'ScreenH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=primary_color,
        spaceBefore=10,
        spaceAfter=4
    )
    cap_style = ParagraphStyle(
        'ScreenCap',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#475569"),
        spaceAfter=10
    )

    story = []
    story.append(Paragraph("PETROCHEMICAL TRADE ANALYTICS", title_style))
    story.append(Paragraph("PROJECT SCREENSHOTS & DEPLOYED SYSTEM DEMONSTRATION", sub_style))

    # Meta Table
    meta_data = [
        [Paragraph("<b>Curriculum Track:</b>", cap_style), Paragraph("HCL Career Shaper COE & NASSCOM FutureSkills Prime", cap_style)],
        [Paragraph("<b>Candidate Name:</b>", cap_style), Paragraph("Anshu Kumar", cap_style)],
        [Paragraph("<b>Scope:</b>", cap_style), Paragraph("Visual Evidence of Working Application & Analytics", cap_style)]
    ]
    meta_t = Table(meta_data, colWidths=[130, 410])
    meta_t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F1F5F9")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(meta_t)
    story.append(Spacer(1, 10))

    for idx, item in enumerate(SCREENSHOT_ITEMS, start=1):
        path = item["path"]
        if not os.path.exists(path):
            continue

        with PILImage.open(path) as img:
            w, h_px = img.size
            aspect = h_px / w

        # Printable width 530 points
        target_w = 530
        target_h = min(270, target_w * aspect)

        img_flow = RLImage(path, width=target_w, height=target_h)

        block = [
            Paragraph(item["title"], h2_style),
            img_flow,
            Spacer(1, 3),
            Paragraph(f"<b>Functional Overview:</b> {item['description']}", cap_style),
            Spacer(1, 10)
        ]
        story.append(KeepTogether(block))

    doc.build(story)
    print(f"Successfully generated Screenshots PDF: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    os.makedirs("submission_files", exist_ok=True)
    generate_screenshots_docx("submission_files/03_Petrochemical_Trade_Analytics_Project_Screenshots.docx")
    generate_screenshots_pdf("submission_files/03_Petrochemical_Trade_Analytics_Project_Screenshots.pdf")
